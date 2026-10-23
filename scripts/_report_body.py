"""Build the latest Word report from frozen results; does not select or trade."""
from pathlib import Path
import json,csv
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'.build/report_base.docx'
OUT.parent.mkdir(parents=True,exist_ok=True)
doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2);sec.bottom_margin=Cm(1.9);sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
sec.footer_distance=Cm(.85)
for name,size in [('Normal',10.5),('Title',23),('Subtitle',12),('Heading 1',17),('Heading 2',12)]:
 s=doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Songti SC')
 s.paragraph_format.space_after=Pt(5);s.paragraph_format.line_spacing=1.18
 if 'Heading' in name:s.paragraph_format.space_before=Pt(8)
normal=doc.styles['Normal'];normal.paragraph_format.widow_control=True
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('A股小市值策略研究  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
doc.core_properties.title='A股小市值策略的收益延续性与执行约束'
doc.core_properties.subject='样本内选组 历史留出验证 风险解释与资金规模敏感性'
doc.core_properties.author=''

def p(t,bold=False,small=False):
 x=doc.add_paragraph();r=x.add_run(t);r.bold=bold
 if small:r.font.size=Pt(9)
 return x

def heading_text(t):return t.replace('、',' ').replace('：',' ').replace('，',' ')
def h(t):doc.add_heading(heading_text(t),2)
def page(t):
 x=doc.add_heading(heading_text(t),1);x.paragraph_format.page_break_before=True
def link(par,label,url):
 x=OxmlElement('w:hyperlink');x.set(qn('r:id'),par.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'245D86');pr.append(c)
 sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');pr.append(sz);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);x.append(r);par._p.append(x)
def table(headers,rows,widths,fs=9.5):
 t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for i,txt in enumerate(headers):t.rows[0].cells[i].text=txt
 for row in rows:
  cells=t.add_row().cells
  for i,txt in enumerate(row):cells[i].text=str(txt)
 for ri,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');trPr.append(ns)
  if ri==0:
   rep=OxmlElement('w:tblHeader');trPr.append(rep)
  for ci,cell in enumerate(row.cells):
   cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EBEE' if ri==0 else ('F5F7F8' if ri%2==0 else 'FFFFFF'));cp.append(sh)
   margins=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'60' if side in ['top','bottom'] else '85');el.set(qn('w:type'),'dxa');margins.append(el)
   cp.append(margins)
   for par in cell.paragraphs:
    par.paragraph_format.space_after=Pt(1);par.paragraph_format.space_before=Pt(1);par.paragraph_format.line_spacing=1.15
    for run in par.runs:run.font.size=Pt(fs);run.bold=ri==0
 borders=OxmlElement('w:tblBorders')
 for s in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+s);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
 t._tbl.tblPr.append(borders)
 gap=doc.add_paragraph();gap.paragraph_format.space_after=Pt(0);gap.paragraph_format.line_spacing=Pt(3);gap.add_run().font.size=Pt(3)
 return t

import pandas as pd,numpy as np,hashlib
R=ROOT/'results';F=R/'word_figures'
def read(name):return pd.read_csv(R/name,float_precision='round_trip')
P=read('performance.csv');M=read('monthly.csv');C=read('cost_sensitivity.csv');T=read('terminal_sensitivity.csv');S=read('st_ablation_performance.csv');G=read('regression_summary.csv');RC=read('regression_coverage.csv');EV=read('events.csv')
SEL=json.load(open(R/'selection.json'));Q=json.load(open(R/'quality.json'));VERIFY=json.load(open(R/'verification.json'));PAY=json.load(open(R/'slide_payload.json'));CFG=json.load(open(ROOT/'config.json'))
B=list(CFG['buckets']);WIN=SEL['winner'];L={s:f'{v[0]}—{v[1]}' for s,v in CFG['buckets'].items()};L['CSI1000']='中证1000'
assert WIN=='R001_100' and len(B)==5 and CFG['main_cost_bps']==30 and VERIFY['errors']==[]
assert CFG['train_start']=='2015-01' and CFG['test_end']=='2025-12'
for s in B:
 x=M[(M.strategy==s)&(M.cost_bps==30)&M.return_month.between('2015-01','2023-12')].net_return
 assert len(x)==108 and np.isclose(np.sqrt(12)*x.mean()/x.std(ddof=1),SEL['scores'][s]['score'])
def met(s,per='TEST'):return P[(P.strategy==s)&(P.period==per)&(P.cost_bps==30)].iloc[0]
def fmt(x,n=2):return '—' if pd.isna(x) else f'{100*float(x):.{n}f}%'
def num(x,n=3):return '—' if pd.isna(x) else f'{float(x):.{n}f}'
def pp(x,n=2):return f'{100*float(x):.{n}f}'
def note(t):p(t,small=True)
def cap(t):
 x=p(t,bold=True);x.paragraph_format.keep_with_next=True

def fig(name,caption,source):
 cap(caption);x=doc.add_paragraph();x.paragraph_format.keep_with_next=True
 run=x.add_run();run.add_picture(str(F/(name+'.png')),width=Cm(16.7))
 inline=run._r.find('.//'+qn('wp:docPr'))
 if inline is not None:inline.set('descr',caption)
 note(source)
def nt(headers,rows,widths,fs=9.2):
 t=table(headers,rows,widths,fs)
 for row in t.rows:
  for cell in row.cells[1:]:
   for par in cell.paragraphs:par.alignment=WD_ALIGN_PARAGRAPH.CENTER
 return t
K=read('supplement/capital_performance.csv');E=read('supplement/extended_performance.csv')
SUP=json.loads((R/'supplement/manifest.json').read_text())
assert SUP['checks']['frozen_selection_unchanged'] and SUP['checks']['missing_holding_marks']==0
used_figures=[]
basefig=fig
def fig(name,caption,source):
 used_figures.append(name);basefig(name,caption,source)
def km(capital,mode='LOT_MIN5',period='TEST'):
 return K[K.initial_capital.eq(capital)&K['mode'].eq(mode)&K.period.eq(period)].iloc[0]
KB=km(1,'CONTINUOUS');PAIR=json.load(open(R/'paired_tests.json'))['comparisons']
CI=next(x for x in PAIR if x['strategy']==WIN and x['comparator']=='CSI1000' and x['period']=='TEST')

# 1 Summary
p('A股小市值策略的收益延续性\n与执行约束').style='Title'
p('2015至2025年沪深主板市值分组研究').style='Subtitle'
note('2026年10月23日  |  样本内2015—2023  |  历史留出2024—2025  |  单边综合成本30bp')
h('研究问题')
p('在剔除当时已知ST等风险状态后，较小市值股票的收益优势能否延续？如果以100只股票为月度目标，这种优势需要承担多大回撤，又会受到整手交易和最低佣金怎样的影响？')
h('主要结论')
p(f'五档组合中，1—100组以样本内夏普{num(met(WIN,"TRAIN").sharpe)}胜出。冻结该组后，其2024—2025年复合年化收益为{fmt(met(WIN).cagr)}，在原五档中最高；但夏普并非最高，日度最大回撤达{fmt(met(WIN).daily_max_drawdown)}。因此可以认可本样本的收益领先，不能推出严格单调或低风险的“市值越小越好”。')
p(f'相对中证1000价格指数，最小组月均收益差为{pp(CI["mean"],3)}个百分点，HAC(6)的95%区间跨越0。ST退出与补仓未明显降低样本外最大回撤；扩展至第1000名后，样本内市值梯度较强，样本外仍有区间反转。')
p(f'资金实验显示，2024年初投入10万元、执行整手及最低佣金规则时，平均现金占比为{fmt(km(100000).mean_cash)}；100万元时降至{fmt(km(1000000).mean_cash)}，对应年化收益{fmt(km(1000000).cagr)}。这是订单约束敏感性，尚不是实盘结算或市场容量结论。')
cap('表1  原策略的主要绩效')
nt(['对象与期间','复合年化收益','夏普比率','日度最大回撤'],[(f'最小组 / {"样本内" if per=="TRAIN" else "样本外"}',fmt(met(WIN,per).cagr),num(met(WIN,per).sharpe),fmt(met(WIN,per).daily_max_drawdown)) for per in ['TRAIN','TEST']]+[('中证1000 / 样本外',fmt(met('CSI1000').cagr),num(met('CSI1000').sharpe),fmt(met('CSI1000').daily_max_drawdown))],[5.1,3.9,3.4,4.4])
note('夏普以月收益计算并年化，无风险收益率设为0。组合扣综合交易成本；中证1000为不含分红、未扣实施费用的价格指数。资金实验另从期初现金建仓，不能与原策略承接持仓的样本外收益直接相减。')
p('2024—2025是历史留出的评价期，研究过程中已查看结果，不能称为未窥视的独立测试。壳价值仅为背景，本研究不识别壳价值或因果关系。',bold=True)

# 2 Protocol
page('1  先固定股票池与投资规则')
p('研究对象为沪深主板，包含原中小板及后来退市公司；排除金融、CDR和上市不足252个交易日的股票。数据以Tushare冻结输入为基础，补以历史简称、已发布行业分类及交易所公告，避免用当前存续名单或当前状态回填历史。')
h('市值名次的含义')
p('每月末先跳过当时已知的ST、PT、退市整理及状态未知股票，再按当时总市值升序排列；同值按代码排序。总市值是总股本对应的市值，不是流通市值。本文“1—100名”始终指过滤后的名次，财报缺失不改变这一排名。')
cap('表2  固定的研究协议')
table(['项目','具体设定'],[
('候选与目标','1—100、101—200、201—300、301—400、401—500；每组100只，月度等权目标。'),
('样本内选优','2015—2023共108个月，按单边30bp成本后的月收益夏普选组。完全并列时先低换手，再较大市值区间。'),
('历史留出验证','冻结赢家与规则后评价2024—2025共24个月；不得用测试排名替换赢家。'),
('主成本与收益','买卖各按实际成交额扣30bp；统一复权收益，不重复加股息。主回测为连续份额。'),
('停牌与退市','无法交易时继续持有；停牌沿用末次估值。退市主残值0，另作末次估值50%和100%情景。'),
],[3.3,13.5],10)
h('从信号到实际成交')
p('月末收盘形成名单；用此前收盘价格和净值确定目标数量。开盘前按已公布涨停价预留买入预算，开盘可交易才执行；停牌或开盘封于相应涨跌停时不成交。普通买入缺口在之后交易日继续尝试，不能依据本次开盘结果临时更换标的或扩大其他订单。')
p('持仓出现可观测ST等风险时，发出退出订单并在首个允许卖出的开盘执行。卖出所得最早用于下一交易日补仓，按当月冻结名单从本组上界之后顺延。受阻退出、零碎余仓和资金等待均可使实际持股数偏离100。')
note('财报公告后下一交易日可用；行业与状态采用历史资料。简称证据最多沿用185天，未知不视为正常股。历史版本与估值限制见第9节。')

# 3 Main results
page('2  从样本内赢家到历史留出表现')
p('五档均参与样本内选优，1—100组的夏普最高。验证阶段保留同一赢家，同时列出其他组，检查优势是否依赖单一指标。')
fig('nav_periods','图1  五档组合与中证1000的月末净值','各期净值归一化为1；样本外承接2023年末实际持仓。组合扣单边30bp，指数为价格指数。')
cap('表3  五档的选择依据与样本外结果')
nt(['过滤后名次','样本内\n夏普','样本外\n年化收益','样本外\n夏普','样本外\n日度最大回撤'],[(L[s],num(met(s,'TRAIN').sharpe),fmt(met(s).cagr),num(met(s).sharpe),fmt(met(s).daily_max_drawdown)) for s in B],[3.4,3.0,3.6,2.9,3.9],10)
p('样本外最小组的年化收益仍领先，但401—500组的夏普为0.906，高于最小组的0.784；301—400与401—500之间也存在收益反转。“最小组本次收益较好”与“所有区间严格单调、最小组在所有维度最优”是不同命题。')
p('样本外两年表现并不均匀：最小组2024年收益约1.05%，2025年约50.74%；中证1000分别约1.20%和27.49%。因此整体年化数字不能替代逐年表现，也不足以判断未来不同市场状态。')
note('2015年最小组收益约158.16%，对较长样本内业绩有明显贡献。样本跨越牛熊阶段并不等于已经获得多个独立市场周期的验证。')

# 4 uncertainty
page('3  收益领先的证据有多强')
fig('paired_intervals','图2  冻结赢家与其他组合的配对月收益差','正值表示最小组领先；横线为HAC(6)标准误构造的95%区间。所有样本外比较均含0。')
p(f'主要统计原假设为H0：最小组与中证1000的总体月均收益差为0。观测差为{pp(CI["mean"],3)}个百分点/月，95%区间为[{pp(CI["ci95_low"],3)}, {pp(CI["ci95_high"],3)}]个百分点/月，因而不能在该检验下拒绝H0。区间包含0不等于证明没有优势，而是现有24个月证据不足以排除无优势的解释。')
note('HAC允许误差存在异方差和序列相关，6为采用的滞后阶数；区间是月均差的区间，不是复合年化差的区间。小样本下推断仍不稳定，且未对全部探索性比较进行多重检验校正。')
fig('test_drawdown','图3  样本外的日度回撤','回撤以各评价期内的历史峰值计算，包含期初净值。最小组日度最大回撤−45.16%，中证1000为−27.08%。')
p('最小组的月末净值最大回撤仅为−33.53%，会遗漏月内更深的下跌。因此正文以日度回撤为主要风险指标，不能把较高收益解释为免费收益或低风险套利。')

# 5 ST
page('4  ST规则和交易成本能否改变判断')
p('为判断结果是否来自风险处置，对同一五档分别运行三套固定规则。以下聚焦冻结赢家；三套规则均在月末过滤已知风险状态，差别仅在持仓期间是否退出、退出后是否补仓。')
cap('表4  最小组的ST规则对照 2024—2025')
rulelabels={'FILTER_ONLY':'仅月度过滤','EXIT_CASH':'触发退出 次月前留现金','EXIT_REPLACE':'触发退出 下一交易日起补仓'}
nt(['规则','年化收益','夏普','日度最大回撤'],[(lab,fmt((v:=S[S.strategy.eq(WIN)&S.period.eq('TEST')&S.rule.eq(rule)] .iloc[0]).cagr),num(v.sharpe),fmt(v.daily_max_drawdown)) for rule,lab in rulelabels.items()],[7.1,3.2,2.8,3.7],10)
p('退出并补仓的年化收益略低于仅月度过滤，三套规则的最大回撤几乎相同。相关配对差的不确定性区间也跨越0，因此不能宣称ST卖出规则已经带来可靠的收益或回撤改善。ST处置仍有风险管理含义，但不应以此忽略尚未披露风险和已发生的价格损失。')
h('有卖出规则不等于能够及时离场')
p('最小组全样本共有128次风险退出成交，其中97次发生延迟。002473.SZ的模型退出信号在2017年5月2日触发，直到12月29日才成交，延迟166个交易日；这些受阻期间的损益仍保留在组合。新增ST、停牌或退市不能事后从历史样本删除。')
cap('表5  最小组的单边成本敏感性')
nt(['单边综合成本','样本内年化收益','样本外年化收益','样本外夏普'],[(f'{cost}bp',fmt((v:=C[C.strategy.eq(WIN)&C.cost_bps.eq(cost)&C.period.eq('TRAIN')].iloc[0]).cagr),fmt((w:=C[C.strategy.eq(WIN)&C.cost_bps.eq(cost)&C.period.eq('TEST')].iloc[0]).cagr),num(w.sharpe)) for cost in [10,30,50]],[3.8,4.4,4.4,4.2],10)
p('成本不仅影响收益水平，也可能改变排序：10bp下，401—500组样本外年化收益高于最小组。退市主残值为0时，最小组样本外年化23.42%；按末次估值50%和100%处理时约为23.58%和23.73%。这只是残值情景，不代表核实后的清算回收款。')

# 6 controls
page('5  已知特征解释了多少区间差异')
p('在每月前500名内，以个股下一月复权毛收益为被解释变量。简化模型只放入四个区间虚拟变量，以1—100组为参照；完整模型加入历史行业、过去20交易日平均成交额的对数、上月收益和归母净利润TTM是否亏损。')
p('两个模型使用完全相同的有效股票样本，连续解释变量按月横截面1%与99%缩尾，再将各月系数按时期汇总并计算HAC(6)区间。回归不用于选股，也不参与更换赢家。')
fig('regression_intervals','图4  控制前后相对最小组的毛收益差','每个点代表月度横截面系数的时间平均；横线为HAC(6) 95%区间。完整模型与简化模型样本一致。')
p('样本内加入控制后，负向区间差异整体缩小，但部分差异仍存在；样本外各区间的估计区间均包含0。结果说明部分差异与这些特征共同变化，既不能证明因果关系，也不能把剩余系数命名为成本后策略Alpha或壳价值。')
cap('表6  控制回归的有效样本覆盖 132个月')
cov=RC.groupby('bucket').coverage.agg(['mean','min'])
nt(['过滤后名次','平均覆盖率','最低单月覆盖率'],[(L[s],fmt(cov.loc[s,'mean']),fmt(cov.loc[s,'min'])) for s in B],[5.6,5.6,5.6],10)
note('缺少财报或控制变量只影响回归样本，不改变投资组合名次。最小组覆盖较弱，因此回归结果描述的是可观测有效样本，不能无条件外推到全部持仓。壳价值若要直接识别，需要借壳事件等额外证据。')

# 7 exactly one page extension
page('6  扩展到第1000名后是否仍有市值梯度')
p('新增501—600、601—700、701—800、801—900、901—1000五档，每档100只，沿用同一历史股票池、ST规则、月度调整和30bp成本。新增区间只作事后扩展检验，不进入原五档选优，不更换冻结赢家。')
fig('extended_bands','图5  从第1名到第1000名的收益与夏普','虚线右侧为新增区间；左侧五档是原选优候选。均为过滤后的总市值排名，不是原始全市场排名。')
cap('表7  新增五档的描述性绩效')
EB=list(E.strategy.drop_duplicates())
nt(['扩展名次','样本内\n年化收益','样本外\n年化收益','样本外\n夏普','样本外\n日度回撤'],[(s[1:].replace('_','—'),fmt(E[E.strategy.eq(s)&E.period.eq('TRAIN')].iloc[0].cagr),fmt((v:=E[E.strategy.eq(s)&E.period.eq('TEST')].iloc[0]).cagr),num(v.sharpe),fmt(v.daily_max_drawdown)) for s in EB],[3.4,3.5,3.5,2.8,3.6],9.7)
p('样本内新增区间的年化收益从约3.94%下降至−2.32%，与较强的小市值收益梯度一致。样本外新增五档为8.12%—16.26%，均低于最小组，但501—800之间反而递增，原五档中也有反转。扩展支持“本样本最小组收益较好”，仍不支持严格单调规律。')
note('新增覆盖为132个月×500只目标股票。000693.SZ在2018-07-13至2019-05-24暂停上市的209个交易日已用原始公告补证，保留受阻持仓及损失；没有补造报价。停牌末次估值仍可能推迟损失确认。本页未新增特征回归或独立样本，不作为确认性显著性证据。')

# 8 capital method
page('7  把资金规模转化为可检验的订单约束')
p('资金实验只针对冻结赢家1—100组。分别在2015年初和2024年初从现金开始，投入10、20、30、50、100、200、500万元，期间收益留在账户，不追加资金。两段实验分别与相同期初现金建仓的连续份额组合比较，避免把前期累计盈利误作初始资金的执行能力。')
cap('表8  分步加入执行约束')
table(['情景','订单数量','成本口径'],[
('连续份额对照','保留原有连续数量，其他时序相同。','单边30bp，无最低收费。'),
('整手','买单向下取100股整数倍；部分减仓向下取整手，清仓释放剩余余额。','单边30bp，无最低收费。'),
('整手加最低佣金','沿用整手约束；含最低收费的预算不足一手时不买入，不替换为别的股票。','27.5bp综合其他成本＋max（2.5bp佣金，5元），买卖分别计费。'),
],[3.3,7.0,6.5],10)
p('2.5bp佣金是显式情景参数，并非已取得的个人账户报价。把它作为30bp综合成本中的一部分，再增加不足5元的差额，可避免重复收取佣金。每个股票、每个交易日、每个买卖方向的一笔成交按一笔订单计费；未成交不收费。')
cap('表9  每笔费用示例 单位为元')
nt(['成交金额','30bp原成本','含最低佣金成本','增加费用'],[(f'{amount:,.0f}',f'{amount*.003:.2f}',f'{amount*.00275+max(amount*.00025,5):.2f}',f'{max(5-amount*.00025,0):.2f}') for amount in [1000,10000,20000]],[4.2,4.2,4.2,4.2],10)
h('数量与资金约束怎样执行')
p('按当日复权因子将连续份额换算为未复权股数，确定100股买入单位；在开盘前以涨停价和完整费用锁定预算。各订单预算总额不得超过前收盘现金，卖款最早下一交易日使用。不足一手的预算留作现金，不用当日价格重新顺延选股。')
h('这项实验的边界')
p('复权账户仍假设分红等收益自动再投资，未逐笔重建派息到账、红股、税款及真实整数持股登记；清仓时可能释放复权记账形成的非整数等价份额。因此这是“订单整手与费用敏感性”，不能标为完整实盘账户模拟。成交额参与率、盘口排队和冲击成本也未加入，500万元情景不等于验证500万元容量。',bold=True)

# 9 capital results
page('8  小资金的主要代价是闲置与费用')
fig('capital_sensitivity','图6  期初资金与年化收益的执行敏感性','两段分别从现金开始；虚线为同起点的连续份额对照。整手加最低佣金是在整手基础上加入最低收费，不是重新选优。')
cap('表10  2024年初投入资金后的结果 2024—2025')
nt(['期初\n万元','仅整手\n年化收益','加最低佣金\n年化收益','加最低佣金\n夏普','平均现金\n占比','平均\n持股数'],[(str(int(capital/10000)),fmt(km(capital,'LOT').cagr),fmt(km(capital).cagr),num(km(capital).sharpe),fmt(km(capital).mean_cash),f'{km(capital).mean_holdings:.1f}') for capital in SUP['protocol']['capitals']],[2.1,2.9,3.4,3.1,2.8,2.5],9.5)
note(f'后3列均为“整手加最低佣金”情景。平均持股数包含未清零余仓和受阻退出，超过100不表示增加选股名额。相同现金起点的连续份额对照年化{fmt(KB.cagr)}、夏普{num(KB.sharpe)}、平均现金{fmt(KB.mean_cash)}；原策略样本外23.42%承接前期持仓，口径不同。')
p(f'10万元时平均仅持有{km(100000).mean_holdings:.1f}只、现金约{fmt(km(100000).mean_cash)}，已经偏离100只股票等权目标。100万元时平均现金约{fmt(km(1000000).mean_cash)}，年化收益较同起点连续份额对照低{pp(KB.cagr-km(1000000).cagr)}个百分点；其中最低佣金相对仅整手情景再降低约{pp(km(1000000,"LOT").cagr-km(1000000).cagr)}个百分点。')
p(f'小资金情景的夏普或回撤较好，部分源于较低仓位：10万元加最低佣金时日度回撤为{fmt(km(100000).daily_max_drawdown)}，100万元时为{fmt(km(1000000).daily_max_drawdown)}。这不是选股能力提升。100万—200万元在本模型中更接近连续份额暴露，但不能据此给出保证收益的资金门槛或可承载资金上限。')

# 10 conclusion
page('9  结论与仍然不能回答的问题')
h('当前证据支持什么')
p('在规定股票池、历史数据、ST和交易成本口径下，最小100只组合是2015—2023年的夏普赢家；冻结规则后，其2024—2025年收益点估计继续领先原五档及新增五档。新增区间加强了对样本内市值梯度的描述，资金实验则说明实现接近研究组合的暴露需要承担整手、现金与费用约束。')
h('哪些原始判断需要收窄')
p('不能把结果写成“最小市值必然最好”。样本外夏普不是最小组最高，各档收益不严格单调，相对基准的月均收益差尚未排除0。ST退出不保证及时成交，也未在样本外显著降低最大回撤。加入控制后仍有差异，不足以识别独立Alpha、壳价值或因果机制。')
h('优先考虑的证据限制')
p('验证独立性：历史留出期已经在研究迭代中查看，新增区间与资金实验都是事后补充；24个月样本较短，且没有全面控制多重比较。下一步应先登记规则，再观察真正未使用的新数据，而不是继续按已见结果调整赢家。')
p('价格与结算：日线开盘能否成交是简化假设，未建模排队、成交量和冲击。停牌沿用末次价格会延迟风险显现，退市残值并非认证清算款；整手实验仍采用复权收益账户。若要给出实际可执行资金需求，应优先补齐公司行动与券商费用结算，并加入成交量限制。')
p('历史信息：早期停牌总市值存在7,879行估计补全，财报历史版本未全部认证，简称快照和季度行业分类也可能修订或滞后。程序检查完整不等于每项数据已经获得独立人工确认。')
p('比较口径：股票使用复权收益，中证1000使用价格指数。未计入指数分红和实际ETF费用，因此本文的相对收益不是可直接交易的净超额收益。')
h('项目的最终表述')
p('本项目验证了一个受ST和交易约束影响的A股小市值策略，并通过历史留出、区间扩展和资金敏感性划定适用边界。核心价值是可复现地说明收益优势、风险代价和证据不足之处，而不是把历史最高收益包装为稳健获利规律。',bold=True)

# 11 appendix
page('附录  数据核验 复现与参考来源')
h('结果与核验记录')
p('主结果覆盖132个月、66,000个分组公司—月，冻结原始输入12,579份。核验清单包含24个公司—月及31个执行案例；源项目核心135项测试通过；公开版本另行执行测试，结果见仓库持续集成记录，新增资金实验核对买单100股整数倍、费用与现金预算，扩展实验补证209个暂停上市交易日。上述为程序与来源记录检查，未获得独立人工签字。')
p('收益：复合年化＝各月（1＋收益率）乘积的12/N次方减1。夏普＝√12×月收益均值÷月收益样本标准差，无风险率0。换手为实际买卖金额之和除以前收盘净值后累积；日度最大回撤包含评价期期初净值。HAC(6)用于月均收益差与平均回归系数的不确定性，不用于CAGR差。',small=True)
h('复现与定位')
p('在项目目录运行单元测试；核心流程为python3 -m microcap all。补充实验运行python3 scripts/run_execution_sensitivity.py；图表与Word分别由scripts/build_word_figures.py、scripts/build_word_report.py生成。Notebook读取冻结结果并独立复算关键指标。依赖、字体和数据授权范围见README。',small=True)
p('核心结果：selection.json、performance.csv、paired_tests.json、st_ablation_*、regression_*。新增结果：results/supplement/内的capital_performance.csv、capital_monthly.csv、extended_performance.csv及交易明细；manifest.json记录输入和输出哈希。公告证据位于data/repair_sources/extension。正文图表来源均可追溯到上述表，Word来源清单为word_report_manifest.json。',small=True)
h('参考来源')
refs=[
('1  Lee Qu Shen 2023 Gate Fees','https://revfin.org/gate-fees-the-pervasive-effect-of-ipo-restrictions-on-chinese-equity-markets/','Review of Finance 27(3) 809–849；壳价值背景，非本文识别结果。'),
('2  Asness等 Size Matters If You Control Your Junk','https://www.aqr.com/insights/research/working-paper/size-matters-if-you-control-your-junk','市值与质量特征控制的方法参考。'),
('3  Harvey Liu Zhu 多重检验与收益因子','https://www.nber.org/papers/w20592','NBER 20592；本文保留多重探索带来的解释限制。'),
('4  深交所 主板交易申报数量说明','https://investor.szse.cn/knowledge/t20230308_599142.html','买入100股整数倍，不足100股余额一次性卖出。'),
('5  深交所 投资者入市手册股票篇','https://www.szse.cn/www/investor/institute/bookshelf/manualseriesbook/P020190322685818724112.pdf','最低5元佣金的情景依据；2.5bp为研究假设。'),
('6  华泽钴镍 关于公司股票暂停上市的公告','https://disc.static.szse.cn/download/disc/disk01/finalpage/2018-07-13/40f49187-2d30-44a2-a031-e8b8d8ec5208.PDF','2018-106；2018年7月13日起暂停上市。'),
('7  华泽钴镍 关于公司股票终止上市的公告','https://static.cninfo.com.cn/finalpage/2019-05-18/1206283352.PDF','2019-029；2019年5月27日进入退市整理期。'),
]
for title,url,desc in refs:
 z=doc.add_paragraph();z.paragraph_format.space_after=Pt(1);link(z,title,url);note(desc)

# Deterministic typography and output provenance.
for par in doc.element.iter(qn('w:p')):
 pr=par.find(qn('w:pPr'))
 if pr is None:pr=OxmlElement('w:pPr');par.insert(0,pr)
 el=OxmlElement('w:snapToGrid');el.set(qn('w:val'),'0');pr.append(el)
for style in doc.styles:
 for el in list(style.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
 for rf in style.element.iter(qn('w:rFonts')):
  for attr in list(rf.attrib):
   if 'theme' in attr.lower():del rf.attrib[attr]
for run in doc.element.iter(qn('w:r')):
 rp=run.find(qn('w:rPr'))
 if rp is None:rp=OxmlElement('w:rPr');run.insert(0,rp)
 rf=rp.find(qn('w:rFonts'))
 if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
 rf.set(qn('w:eastAsia'),'Songti SC')
doc.save(OUT)
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
inputs=['config.json','README.md','results/selection.json','results/performance.csv','results/monthly.csv','results/daily_nav.csv','results/paired_tests.json','results/cost_sensitivity.csv','results/terminal_sensitivity.csv','results/st_ablation_performance.csv','results/st_ablation_tests.csv','results/regression_summary.csv','results/regression_coverage.csv','results/verification.json','scripts/_report_body.py','scripts/build_word_figures.py','results/supplement/manifest.json','results/supplement/capital_performance.csv','results/supplement/extended_performance.csv']
manifest=dict(report_date='2026-10-23',selection_winner=WIN,cost_bps=30,source_inputs=[dict(path=x,sha256=sha(ROOT/x)) for x in inputs],output=dict(path=str(OUT.relative_to(ROOT)),sha256=sha(OUT)),figures=[dict(path=str((F/(name+'.png')).relative_to(ROOT)),sha256=sha(F/(name+'.png'))) for name in used_figures],validation=dict(training_scores_recomputed=True,existing_release_errors=VERIFY['errors'],render_review='pending',tables=len(doc.tables),figures=len(doc.inline_shapes)))
(R/'word_report_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(OUT);print('Tables',len(doc.tables),'figures',len(doc.inline_shapes))
