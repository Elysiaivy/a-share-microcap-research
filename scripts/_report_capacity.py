"""Insert the completed capacity study into the latest report.

Call after _report_applicability.py and build_capacity_figures.py. This module's
data validator uses only the standard library so the plotting script can reuse
it without loading python-docx. Actual document authoring must use the bundled
artifact runtime; this script does not render or certify the output layout.
"""
from pathlib import Path
import argparse
import csv
import hashlib
import itertools
import json
import math

ROOT = Path(__file__).resolve().parents[1]
REQUIRED_SCENARIOS = {
    'OPEN_BASE', 'VWAP_BASE', 'CAP_1P', 'CAP_1P_SIGMA', 'MAIN_1P',
    'TIGHT_0P5', 'LOOSE_5P', 'IMPACT_HIGH_1P',
}
COMPARISONS = [
    ('VWAP_BASE', 'OPEN_BASE', '价格参照与成交资格'),
    ('CAP_1P', 'VWAP_BASE', '容量与延迟路径'),
    ('CAP_1P_SIGMA', 'CAP_1P', '波动率数据覆盖'),
    ('MAIN_1P', 'CAP_1P_SIGMA', '增量冲击与持仓路径'),
]


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def number(row, key):
    value = float(row[key])
    if not math.isfinite(value):
        raise ValueError(f'Non-finite published metric: {key}={row[key]}')
    return value


def pct(value):
    return f'{float(value):.2%}'


def pp(value):
    return f'{100 * float(value):+.2f}'


def capital_label(value):
    return f'{int(value) / 10000:,.0f}万元'


def load_capacity(root=ROOT):
    """Require the entire frozen grid and independently validate saved outputs."""
    root = Path(root)
    folder = root / 'results/capacity'
    protocol = json.loads((folder / 'protocol.json').read_text())
    manifest = json.loads((folder / 'manifest.json').read_text())
    if manifest.get('full_predeclared_grid') is not True:
        raise ValueError('Only a completed full predeclared grid can enter the report.')
    if set(protocol['scenarios']) != REQUIRED_SCENARIOS:
        raise ValueError('Capacity report requires all eight predeclared scenarios.')
    if protocol['periods'] != {'TRAIN': ['2015-01', '2023-12'], 'TEST': ['2024-01', '2025-12']}:
        raise ValueError('Unexpected evaluation periods.')
    if protocol.get('winner') != 'R001_100' or protocol.get('base_cost_bps') != 30:
        raise ValueError('Capacity study must preserve the frozen winner and cost.')
    if protocol['scenarios']['CAP_1P_SIGMA'].get('require_known_sigma') is not True:
        raise ValueError('Missing matched volatility-coverage control.')
    expected_capitals = [100000, 300000, 1000000, 2000000, 5000000, 10000000, 30000000, 50000000]
    if sorted(protocol['capitals_yuan']) != expected_capitals:
        raise ValueError('Report wording requires the frozen eight capital levels.')
    scenario_rules = {'OPEN_BASE': (None, 0), 'VWAP_BASE': (None, 0),
                      'CAP_1P': (.01, 0), 'CAP_1P_SIGMA': (.01, 0),
                      'MAIN_1P': (.01, .5), 'TIGHT_0P5': (.005, .5),
                      'LOOSE_5P': (.05, .5), 'IMPACT_HIGH_1P': (.01, 1.)}
    for name, (participation, eta) in scenario_rules.items():
        if (protocol['scenarios'][name]['participation'], protocol['scenarios'][name]['impact_eta']) != (participation, eta):
            raise ValueError(f'Unexpected predeclared parameters: {name}')
    if (protocol['buy_lot_shares'], protocol['minimum_commission_yuan'], protocol['commission_bps']) != (100, 5, 2.5):
        raise ValueError('Unexpected lot or minimum-commission protocol.')
    digest = manifest.get('inputs', {}).get('results/capacity/protocol.json')
    if digest != sha(folder / 'protocol.json'):
        raise ValueError('Protocol changed after the recorded capacity run.')
    tables = {}
    for name in ['performance', 'performance_differences', 'attempt_summary', 'constraint_summary']:
        path = folder / f'{name}.csv'
        if manifest.get('outputs', {}).get(path.name) != sha(path):
            raise ValueError(f'Missing or mismatched output checksum: {path.name}')
        with path.open(newline='', encoding='utf-8-sig') as handle:
            tables[name] = list(csv.DictReader(handle))
        if not tables[name]:
            raise ValueError(f'Empty required capacity output: {path.name}')
    coverage = json.loads((folder / 'coverage.json').read_text())
    if not isinstance(coverage, dict) or not coverage:
        raise ValueError('A nonempty public coverage audit is required.')
    risk_manifest = json.loads((folder/'risk_diagnostics_manifest.json').read_text())
    if (risk_manifest['run_id'] != manifest['run_id']
            or risk_manifest['capacity_manifest_sha256'] != sha(folder/'manifest.json')
            or risk_manifest['analysis_script_sha256'] != sha(root/'scripts/analyze_capacity_risk.py')):
        raise ValueError('Risk diagnostics do not match the frozen capacity run.')
    for name in ['risk_diagnostics','terminal_cases']:
        path = folder/f'{name}.csv'
        if risk_manifest['outputs'].get(path.name) != sha(path):
            raise ValueError('Risk diagnostics output checksum differs.')
        with path.open(newline='',encoding='utf-8-sig') as handle:
            tables[name] = list(csv.DictReader(handle))
    expected = set(itertools.product(protocol['periods'], protocol['scenarios'], protocol['capitals_yuan']))
    perf = tables['performance']
    keys = [(r['period'], r['scenario'], int(r['initial_capital'])) for r in perf]
    if len(keys) != len(set(keys)) or set(keys) != expected:
        raise ValueError('Performance table does not contain the full unique protocol grid.')
    if manifest.get('n_scenarios') != len(expected):
        raise ValueError('Manifest job count does not match the full grid.')
    indexed = dict(zip(keys, perf))
    risk_index = {(r['period'],r['scenario'],int(r['initial_capital'])):r for r in tables['risk_diagnostics']}
    if len(risk_index) != len(tables['risk_diagnostics']) or set(risk_index) != expected:
        raise ValueError('Incomplete or duplicate terminal-accounting grid.')
    required = ['cagr', 'sharpe', 'daily_max_drawdown', 'mean_cash', 'max_nav_yuan',
                'attempt_completion', 'n_attempts', 'impact_cost_yuan', 'total_traded_yuan',
                'final_pending_risk_stocks', 'final_pending_risk_weight', 'max_risk_sell_delay']
    for key, row in indexed.items():
        if int(row['n_months']) != (108 if key[0] == 'TRAIN' else 24):
            raise ValueError(f'Incomplete return months: {key}')
        for field in required:
            number(row, field)
        if not 0 <= number(row, 'attempt_completion') <= 1 + 1e-10:
            raise ValueError(f'Invalid attempt completion: {key}')
    difference_index = {}
    for row in tables['performance_differences']:
        key = (row['period'], int(row['initial_capital']), row['scenario'], row['reference'])
        if key in difference_index:
            raise ValueError('Duplicate performance difference.')
        a = indexed[(key[0], key[2], key[1])]
        b = indexed[(key[0], key[3], key[1])]
        for metric in ['cagr', 'sharpe', 'daily_max_drawdown', 'mean_cash', 'attempt_completion']:
            actual = number(a, metric) - number(b, metric)
            if not math.isclose(actual, number(row, metric + '_difference'), abs_tol=1e-11):
                raise ValueError(f'Inconsistent saved difference: {key}/{metric}')
        difference_index[key] = row
    expected_diff = {(p, c, a, b) for p in protocol['periods']
                     for c in protocol['capitals_yuan'] for a, b, _ in COMPARISONS}
    if set(difference_index) != expected_diff:
        raise ValueError('Incomplete predeclared comparison grid.')
    attempt_index = {}
    for row in tables['attempt_summary']:
        key = (row['period'], row['scenario'], int(row['initial_capital']), row['side'], row['order_reason'])
        if key in attempt_index:
            raise ValueError('Duplicate order-attempt summary.')
        attempt_index[key] = row
    for key, row in indexed.items():
        summary = attempt_index[(*key, 'all', 'all')]
        planned = number(summary, 'attempt_requested_value')
        if planned <= 0 or not math.isclose(number(summary, 'attempt_filled_value') / planned,
                                             number(row, 'attempt_completion'), abs_tol=1e-11):
            raise ValueError(f'Order-attempt completion does not reconcile: {key}')
    return dict(protocol=protocol, manifest=manifest, coverage=coverage, indexed=indexed, risk_index=risk_index,
                differences=difference_index, attempts=attempt_index, **tables)


def constraint_count(data, period, capital, reasons, scenario='MAIN_1P'):
    return sum(int(r['n_attempts']) for r in data['constraint_summary']
               if r['period'] == period and r['scenario'] == scenario
               and int(r['initial_capital']) == capital and r['constraint_reason'] in reasons)


def coverage_sentence(coverage):
    """Use auditable aggregate counts when the source schema provides them."""
    if 'coverage' in coverage and isinstance(coverage['coverage'], dict):
        coverage = coverage['coverage']
    labels = [('old_new_differences', '新旧重叠字段差异'),
              ('vwap_anomalies_market', '原市场矩阵内VWAP异常'),
              ('missing_volume_on_quoted_market_rows', '已有开盘报价但成交量缺失的公司日')]
    parts = [f'{label}{int(coverage[key]):,}项' for key, label in labels if key in coverage]
    text = '公开输入核验记录：' + '；'.join(parts) + '。' if parts else ''
    plans = coverage.get('order_plans.csv', {})
    if isinstance(plans, dict) and all(k in plans for k in ['n', 'adv_unknown', 'adv_zero', 'sigma_unknown']):
        text += (f'原冻结赢家的{int(plans["n"]):,}条计划记录中，ADV未知{int(plans["adv_unknown"]):,}条、'
                 f'ADV为0有{int(plans["adv_zero"]):,}条、σ未知{int(plans["sigma_unknown"]):,}条。'
                 '该口径是原订单覆盖，不是新场景的成交率。')
    return text or '输入覆盖详情见results/capacity/coverage.json；实际阻单按各场景的尝试记录另行统计。'


def append_capacity_report(root=ROOT, output=None):
    # Defer artifact dependencies so the pure validator is reusable by plotting.
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    root = Path(root)
    data = load_capacity(root)
    target = Path(output) if output else root / 'docs/research_report.docx'
    figure = root / 'results/word_figures/capacity_stress.png'
    if not figure.is_file():
        raise FileNotFoundError('Run build_capacity_figures.py before authoring the report.')
    figure_manifest = json.loads((root / 'results/capacity/figure_manifest.json').read_text())
    if (figure_manifest.get('capacity_run_id') != data['manifest']['run_id'] or
            figure_manifest.get('outputs', {}).get(figure.name) != sha(figure)):
        raise ValueError('Capacity figure does not match the completed result grid.')
    doc = Document(target)
    if any(p.text.startswith('12  成交容量') for p in doc.paragraphs):
        raise ValueError('Regenerate the base report before inserting capacity sections again.')
    anchor = next(p._p for p in doc.paragraphs if p.text.startswith('12  结论'))
    base_hash = sha(target)

    def move(element):
        anchor.addprevious(element)

    def p(text, heading=False, small=False, bold=False):
        par = doc.add_paragraph(text, style='Heading 2' if heading else 'Normal')
        move(par._p)
        par.paragraph_format.space_after = Pt(4)
        par.paragraph_format.line_spacing = Pt(18 if heading else (13 if small else 14))
        for run in par.runs:
            run.font.name = 'Arial'
            run.font.size = Pt(12 if heading else (9 if small else 10))
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.bold = heading or bold
            run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Songti SC')
        return par

    def page(text):
        par = doc.add_paragraph(text, style='Heading 1')
        par.paragraph_format.page_break_before = True
        par.paragraph_format.line_spacing = Pt(24)
        move(par._p)

    def table(title, headers, rows, widths):
        caption = p(title, bold=True)
        caption.paragraph_format.keep_with_next = True
        tab = doc.add_table(rows=1, cols=len(headers))
        tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        tab.autofit = False
        for col, width in zip(tab.columns, widths):
            col.width = Cm(width)
        for cell, text in zip(tab.rows[0].cells, headers):
            cell.text = text
        for values in rows:
            for cell, value in zip(tab.add_row().cells, values):
                cell.text = str(value)
        for i, row in enumerate(tab.rows):
            props = row._tr.get_or_add_trPr()
            props.append(OxmlElement('w:cantSplit'))
            if i == 0:
                props.append(OxmlElement('w:tblHeader'))
            for j, cell in enumerate(row.cells):
                cell.width = Cm(widths[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell_props = cell._tc.get_or_add_tcPr()
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E7EBEE' if i == 0 else ('F5F7F8' if i % 2 == 0 else 'FFFFFF'))
                cell_props.append(shade)
                margins = OxmlElement('w:tcMar')
                for side in ['top', 'bottom', 'left', 'right']:
                    element = OxmlElement('w:' + side)
                    element.set(qn('w:w'), '65')
                    element.set(qn('w:type'), 'dxa')
                    margins.append(element)
                cell_props.append(margins)
                for par in cell.paragraphs:
                    par.paragraph_format.space_before = Pt(1)
                    par.paragraph_format.space_after = Pt(1)
                    par.paragraph_format.line_spacing = Pt(13)
                    par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in par.runs:
                        run.font.name = 'Arial'
                        run.font.size = Pt(9)
                        run.bold = i == 0
                        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Songti SC')
        borders = OxmlElement('w:tblBorders')
        for side in ['top', 'bottom', 'left', 'right', 'insideH', 'insideV']:
            element = OxmlElement('w:' + side)
            for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                element.set(qn('w:' + key), value)
            borders.append(element)
        tab._tbl.tblPr.append(borders)
        move(tab._tbl)
        gap = p('')
        gap.paragraph_format.line_spacing = Pt(3)
        gap.paragraph_format.space_after = Pt(0)

    def metric(period, scenario, capital):
        return data['indexed'][(period, scenario, capital)]

    def difference(period, capital, a, b):
        return number(data['differences'][(period, capital, a, b)], 'cagr_difference')

    def math_run(text):
        run = OxmlElement('m:r')
        content = OxmlElement('m:t')
        content.text = text
        run.append(content)
        return run

    def impact_equation():
        # Editable OMML, with a true radical and fraction rather than raw LaTeX.
        par = p('')
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        equation = OxmlElement('m:oMath')
        equation.append(math_run('C(Q) = η × σ × Q × '))
        rad = OxmlElement('m:rad')
        rad_pr = OxmlElement('m:radPr')
        hidden = OxmlElement('m:degHide')
        hidden.set(qn('m:val'), '1')
        rad_pr.append(hidden)
        rad.append(rad_pr)
        rad.append(OxmlElement('m:deg'))
        body = OxmlElement('m:e')
        fraction = OxmlElement('m:f')
        numerator, denominator = OxmlElement('m:num'), OxmlElement('m:den')
        numerator.append(math_run('Q'))
        denominator.append(math_run('ADV'))
        fraction.extend([numerator, denominator])
        body.append(fraction)
        rad.append(body)
        equation.append(rad)
        par._p.append(equation)
        par.paragraph_format.line_spacing = 1.3

    page('12  成交容量与增量冲击的固定检验')
    p('本次固定1—100名赢家，增加8档资金、8种执行场景，并分别重跑2015—2023与2024—2025，共128格。'
      '两段均从现金开始，此后不追加或抽走资金。协议在场景结果评估前固定；2024—2025仍是已查看的历史留出，不是新的独立验证。')
    table('表17  预先固定的执行场景', ['场景', '成交参照', '日参与率上限', '增量冲击η'], [
        ['OPEN_BASE 原开盘基准', '次日开盘', '无', '0'],
        ['VWAP_BASE 价格参照', '日VWAP近似', '无', '0'],
        ['CAP_1P 容量对照', '日VWAP近似', '1%', '0'],
        ['CAP_1P_SIGMA 覆盖对照', '日VWAP近似', '1%', '0'],
        ['MAIN_1P 主压力', '日VWAP近似', '1%', '0.5'],
        ['TIGHT_0P5 更紧容量', '日VWAP近似', '0.5%', '0.5'],
        ['LOOSE_5P 更宽容量', '日VWAP近似', '5%', '0.5'],
        ['IMPACT_HIGH_1P 更高冲击', '日VWAP近似', '1%', '1.0'],
    ], [6.4, 3.6, 3.6, 3.2])
    p('所有场景保留100股整手和最低5元佣金。CAP_1P_SIGMA与主压力场景要求同样的已知σ，'
      '但不收增量冲击，用于把σ缺失覆盖变化与冲击路径分开。该对照在输入覆盖核验后、场景结果评估前加入。', small=True)
    p('只把当日数据用于成交结果', heading=True)
    p('选股、目标与预算使用事前信息。ADV是前20个交易日成交额均值，σ是前20个交易日复权日收益的样本标准差，均滞后一日且要求20个已知观察。'
      '当日金额与成交量换算的全天VWAP仅作成交参照，并核对最高最低价；保留原开盘可交易条件，再增加日内停牌、数据有效性和一字板筛选。'
      '因此VWAP与开盘基准之差同时包含价格和资格变化。')
    p('同一股票同一天的买卖共同使用金额上限min（p×滞后ADV，p×当日成交额），而非每笔独享容量。'
      '仅有证据的全天停牌可以令成交额记0；未知值保留未知并阻单。')
    p('费用 预算与部分成交', heading=True)
    impact_equation()
    p('Q为同股同日累计成交金额，每笔只收C（成交后Q）减C（成交前Q）。基础费用为27.5bp其他成本加2.5bp佣金，'
      '佣金每笔最低5元；冲击在此之外另计，不再改价重复收费。η为未校准的压力参数，不是实际冲击成本估计。')
    p('买入按已知涨停价及保守冲击预算安排，仅用前收盘现金；当日卖款不能资助当日买单。普通订单部分成交后继续尝试；'
      'ST剩余卖单及顺延补仓共用原风险事件的一个槽位，后续卖款与未完补仓保留，月末重新设定目标另记日志。')

    page('13  历史留出期的资金规模压力结果')
    p('表18使用相同现金起点逐档比较，基准本身也执行整手与最低佣金。这里的资金是期初投入额，账户随后复利；'
      '不能与原策略承接2023年持仓的23.42%年化直接相减。')
    rows = []
    for capital in data['protocol']['capitals_yuan']:
        a, b = metric('TEST', 'OPEN_BASE', capital), metric('TEST', 'MAIN_1P', capital)
        rows.append([capital_label(capital), pct(number(a, 'cagr')), pct(number(b, 'cagr')),
                     f'{number(b, "sharpe"):.2f}',
                     pct(number(b, 'daily_max_drawdown')), pct(number(b, 'attempt_completion')),
                     pct(number(b, 'mean_cash'))])
    table('表18  开盘基准与主压力场景 2024—2025',
          ['初始资金', '基准年化', '主压力\n年化', '主压力\n夏普', '主压力\n日度回撤', '尝试金额\n完成率', '平均现金'],
          rows, [2.4, 2.4, 2.4, 1.5, 2.7, 2.7, 2.7])
    p('完成率＝各次实际成交份额×前收盘标记价格之和，除以计划份额×同一标记价格之和。'
      '次日重试再次进入分母；该指标衡量每日尝试金额完成程度，不是唯一母单最终成交率。', small=True)
    caption = p('图7  资金增加后的收益与尝试完成率', bold=True)
    caption.paragraph_format.keep_with_next = True
    picture = p('')
    picture.paragraph_format.keep_with_next = True
    picture.paragraph_format.line_spacing = 1.0
    picture.add_run().add_picture(str(figure), width=Cm(16.7))
    p('蓝线为1%主压力，深蓝为0.5%容量，浅蓝为5%容量，三者η均为0.5；灰虚线为原开盘基准。'
      '横轴为期初资金，对数刻度。未依据结果寻找最佳上限。', small=True)
    small, large = metric('TEST', 'MAIN_1P', 1000000), metric('TEST', 'MAIN_1P', 50000000)
    small_base, large_base = metric('TEST', 'OPEN_BASE', 1000000), metric('TEST', 'OPEN_BASE', 50000000)
    p(f'100万元主压力年化{pct(number(small, "cagr"))}，较同资金开盘基准变动'
      f'{pp(number(small, "cagr") - number(small_base, "cagr"))}个百分点，尝试完成率{pct(number(small, "attempt_completion"))}；'
      f'5000万元对应{pct(number(large, "cagr"))}、{pp(number(large, "cagr") - number(large_base, "cagr"))}个百分点和'
      f'{pct(number(large, "attempt_completion"))}。这些是固定假设下的历史压力结果，不能认证实盘可承载资金。')

    page('14  差额来源 数据覆盖与未完成的退出')
    table('表19  分步对照的年化收益变化 2024—2025',
          ['比较及含义', '100万元\n变化百分点', '5000万元\n变化百分点'],
          [[f'{a} − {b}\n{label}', pp(difference('TEST', 1000000, a, b)),
            pp(difference('TEST', 50000000, a, b))] for a, b, label in COMPARISONS], [10.0, 3.4, 3.4])
    p('各行是重跑后的年化差额，不是现金费用的简单扣减。最后一行使用相同σ资格；'
      '差异仍可通过现金、整手、延迟与后续持仓传播。容量更紧可能减少涨跌两侧暴露，曲线不保证单调。', small=True)
    high_small, high_large = metric('TEST', 'IMPACT_HIGH_1P', 1000000), metric('TEST', 'IMPACT_HIGH_1P', 50000000)
    p(f'保持1%容量，把η从0.5升至1.0，100万／5000万元年化分别较主压力变动'
      f'{pp(number(high_small,"cagr") - number(small,"cagr"))}／'
      f'{pp(number(high_large,"cagr") - number(large,"cagr"))}个百分点。主压力这两档累计增量冲击费用为'
      f'{number(small,"impact_cost_yuan") / 10000:,.2f}／{number(large,"impact_cost_yuan") / 10000:,.2f}万元。', small=True)
    train_rows = []
    for capital in [1000000, 50000000]:
        base, main = metric('TRAIN', 'OPEN_BASE', capital), metric('TRAIN', 'MAIN_1P', capital)
        train_rows.append([capital_label(capital), pct(number(base, 'cagr')), pct(number(main, 'cagr')),
                           pct(number(main, 'attempt_completion')),
                           f'{number(main, "max_nav_yuan") / 10000:,.1f}万元'])
    table('表20  样本内压力结果及复利规模 2015—2023',
          ['初始资金', '基准年化', '主压力年化', '尝试完成率', '主压力最大净值'],
          train_rows, [2.7, 3.0, 3.1, 3.4, 4.6])
    p(f'2024—2025主压力的最大净值：100万元起点为{number(small, "max_nav_yuan") / 10000:,.1f}万元，'
      f'5000万元起点为{number(large, "max_nav_yuan") / 10000:,.1f}万元。'
      '九年样本内路径可能放大实际资金量；所有结果均不代表全程维持恒定规模。完整八场景收益、月度夏普和风险指标见performance.csv。', small=True)
    if all('month_end_mean_target_shortfall_weight' in row for row in [small, large]):
        p(f'历史留出100万／5000万元主压力的月末平均普通目标缺口占净值'
          f'{pct(number(small,"month_end_mean_target_shortfall_weight"))}／'
          f'{pct(number(large,"month_end_mean_target_shortfall_weight"))}；'
          '这只度量普通目标缺口，ST补仓未用预算在结果表中另列，不能把尝试完成率当作组合目标完成率。', small=True)
    risk_heading = p('14（续） 未知数据与风险退出需要单列', heading=True)
    risk_heading.paragraph_format.page_break_before = True
    p(coverage_sentence(data['coverage']), small=True)
    blocks = {'missing_or_zero_lag_liquidity', 'missing_or_invalid_daily_liquidity'}
    counts = [constraint_count(data, period, cap, blocks) for period, cap in
              [('TRAIN', 1000000), ('TRAIN', 50000000), ('TEST', 1000000), ('TEST', 50000000)]]
    p(f'主压力中，因所需数据缺失、无效或ADV为0而未成交的尝试，样本内100万／5000万元分别为'
      f'{counts[0]:,}／{counts[1]:,}次，历史留出为{counts[2]:,}／{counts[3]:,}次。'
      '该汇总含零ADV，不可全部称为数据缺失；详细原因见constraint_summary.csv。', small=True)
    p(f'历史留出100万／5000万元主压力的ST卖出已成交部分，最长延迟分别为'
      f'{int(number(small, "max_risk_sell_delay"))}／{int(number(large, "max_risk_sell_delay"))}个交易日；'
      f'期末待退出股票为{int(number(small, "final_pending_risk_stocks"))}／{int(number(large, "final_pending_risk_stocks"))}只，'
      f'占净值{pct(number(small, "final_pending_risk_weight"))}／{pct(number(large, "final_pending_risk_weight"))}。'
      '已成交延迟不覆盖仍未卖出的尾部风险。', small=True)
    if all('final_pending_ST_age_max_sessions' in row for row in [small, large]):
        p(f'期末仍待退出头寸的最长已等待时间分别为{int(number(small, "final_pending_ST_age_max_sessions"))}／'
          f'{int(number(large, "final_pending_ST_age_max_sessions"))}个交易日。', small=True)
    risk_small = data['risk_index'][('TEST','MAIN_1P',1000000)]
    risk_large = data['risk_index'][('TEST','MAIN_1P',50000000)]
    p(f'期末待退出为0不等于全部成功卖出。600321.SH在模型中于2024-06-27按冻结的零残值规则终止记账；'
      f'100万／5000万元主情景的最后账面冲销分别为{number(risk_small,"terminal_writeoff_yuan"):,.0f}／'
      f'{number(risk_large,"terminal_writeoff_yuan"):,.0f}元，卖出回款均为0。'
      '这是终止时账面冲销，非总投资损失；多情景中的同一事件不能重复算作多家公司。', small=True)
    p('日频总量无法描述分钟价格路径、买卖价差、开盘竞价队列及成交量的内生反应。全天VWAP也不保证真实可获得。'
      '本文没有实测订单校准η，复权份额与分红再投资仍不是完整券商结算账户。'
      '因此只判断固定压力假设怎样改变收益与暴露，不宣布“某资金规模一定可执行”。', small=True)

    # Repair the earlier limitations locally and integrate the new evidence.
    for par in doc.paragraphs:
        text = par.text
        if text.startswith('12  结论'):
            par.text = text.replace('12  ', '15  ', 1)
        elif '历史版本与估值限制见第12节' in text:
            par.text = text.replace('见第12节', '见第15节')
        elif text.startswith('资金实验显示，'):
            par.text = (f'资金与执行实验表明，同从2024年初现金开始，100万元在1%日容量及η=0.5增量冲击假设下年化'
                        f'{pct(number(small, "cagr"))}，较同资金开盘基准变动{pp(number(small, "cagr")-number(small_base,"cagr"))}个百分点；'
                        f'5000万元时为{pct(number(large,"cagr"))}，尝试完成率{pct(number(large,"attempt_completion"))}。'
                        '这是固定日频假设下的压力检验，不能认证实盘容量，详见第12—14节。')
        elif text.startswith('复权账户仍假设分红'):
            par.text = (text.split('成交额参与率、')[0] +
                        '本节仅检验整手及最低佣金；第12—14节另加入日成交额上限和增量冲击。盘口排队与实测冲击校准仍未完成。')
        elif text.startswith('交易适用性可以分为成本承受'):
            par.text = ('交易适用性包括成本承受、风险退出、资金分散和市场容量。本节汇总第4及第7—8节的既有实验；'
                        '第12—14节进一步比较同现金起点的成交容量和增量冲击压力。')
        elif text.startswith('本策略更适合被理解为高分散'):
            par.text = ('本策略是高分散并承担较大权益回撤的小市值研究组合。有限资金下，100只目标的整手及费用约束会改变实际暴露。'
                        '资金增加可以减少整手闲置，却可能增加成交容量压力。第12—14节按固定日频上限和增量冲击假设检验这一权衡；'
                        '缺少分钟队列与实测校准，仍不给出认证的实盘资金上限。')
        elif text.startswith('价格与结算：'):
            par.text = ('价格与结算：本次增加了日成交额约束、部分成交重试与增量冲击情景，但没有分钟路径、买卖价差、队列或实测冲击校准，'
                        '日VWAP也不是可保证的成交价。停牌末次估值可延迟风险显现，退市残值并非认证清算款；'
                        '整手实验沿用复权账户，派息到账、红股、税款与整数登记仍需完整重建。')
        elif text.startswith('本项目验证了一个受ST'):
            par.text = ('本项目研究了一个受ST和交易约束影响的A股小市值策略，并用历史留出、区间扩展、收益集中度、市场环境、'
                        '整手费用与成交容量压力划定适用边界。固定压力假设会改变实现收益与持仓路径，不能据此认证可执行资金上限。'
                        '研究支持描述历史收益优势及风险代价，不支持把最高历史收益包装为稳健获利规律、独立Alpha或壳价值。')
        elif text.startswith('核心结果位于results，'):
            par.text = (text + ' 成交容量协议、覆盖记录、128格绩效及尝试统计位于results/capacity；'
                        '原始日行情与逐笔模拟账本仍不公开，结果复算不等于独立认证输入。')
        elif text.startswith('主结果覆盖132个月'):
            par.text = (text + ' 新增128格容量情景分别重建并核对每日现金、持仓净值、前日预算、整手、共享成交额度与费用；'
                        '逐格检查摘要见results/capacity/ledger_checks.csv。这是模型内部一致性验证，不是实际成交认证。')
        elif text.startswith('公开包可用research.ipynb复算'):
            par.text = (text + ' 容量完整重跑使用scripts/run_capacity_stress.py，需授权数据；'
                        'scripts/build_capacity_figures.py复算图表，_report_capacity.py在统一入口中插入本次结果。')

    # Add the methodology sources to the existing references rather than filling
    # the analytical pages with long URLs. External links remain clickable.
    references = [
        ('8  Tushare 日线行情字段说明', data['protocol']['references'][0],
         '成交额以千元、成交量以手记录；用于日VWAP和成交容量输入的单位换算。'),
        ('9  Tóth等 Anomalous Price Impact and the Critical Nature of Liquidity', data['protocol']['references'][1],
         'Physical Review X 1 021006；平方根冲击形式的方法背景，不提供本策略的实测参数。'),
        ('10  Frazzini Israel Moskowitz Trading Costs of Asset Pricing Anomalies', data['protocol']['references'][2],
         'AQR工作论文；交易成本与规模影响的方法参考，不能把海外估计直接当作A股参数。'),
    ]
    for title, url, description in references:
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(1)
        par.paragraph_format.line_spacing = Pt(13)
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), par.part.relate_to(url, RT.HYPERLINK, is_external=True))
        run, props, content = OxmlElement('w:r'), OxmlElement('w:rPr'), OxmlElement('w:t')
        color, size = OxmlElement('w:color'), OxmlElement('w:sz')
        color.set(qn('w:val'), '245D86')
        size.set(qn('w:val'), '18')
        props.extend([color, size])
        content.text = title
        run.extend([props, content])
        link.append(run)
        par._p.append(link)
        note = doc.add_paragraph(description)
        note.paragraph_format.line_spacing = Pt(13)
        note.paragraph_format.space_after = Pt(3)
        for r in note.runs:
            r.font.size = Pt(9)
    for par in doc.paragraphs:
        if par.style.name.startswith('Heading'):
            for run in par.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
        for run in par.runs:
            props = run._element.get_or_add_rPr()
            fonts = props.rFonts
            if fonts is None:
                fonts = OxmlElement('w:rFonts')
                props.insert(0, fonts)
            fonts.set(qn('w:eastAsia'), 'Songti SC')
    doc.core_properties.author = 'Elysiaivy'
    doc.core_properties.last_modified_by = 'Elysiaivy'
    doc.core_properties.subject = '历史留出验证 资金与成交容量压力 风险及证据边界'
    from datetime import datetime, timezone
    report_date = datetime(2026, 10, 25, tzinfo=timezone.utc)
    doc.core_properties.created = report_date
    doc.core_properties.modified = report_date
    doc.save(target)
    provenance = {'base_sha256': base_hash, 'output_sha256': sha(target),
                  'capacity_run_id': data['manifest']['run_id'], 'full_predeclared_grid': True,
                  'inserted_sections': [12, 13, 14], 'conclusion_section': 15,
                  'new_tables': [17, 18, 19, 20], 'new_figures': [7],
                  'render_review': 'pending',
                  'sources': {name: sha(root / 'results/capacity' / name)
                              for name in ['protocol.json', 'performance.csv', 'performance_differences.csv',
                                           'attempt_summary.csv', 'constraint_summary.csv', 'coverage.json', 'manifest.json']},
                  'figure_sha256': sha(figure)}
    (root / 'results/capacity/document_manifest.json').write_text(
        json.dumps(provenance, ensure_ascii=False, indent=2) + '\n')
    print(target)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, default=ROOT)
    parser.add_argument('--output', type=Path)
    args = parser.parse_args()
    append_capacity_report(args.root, args.output)
